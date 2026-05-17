"""
Generate Data Dictionary untuk df_final_full
Output: data_dictionary.xlsx
Jalankan: python generate_data_dictionary.py
"""

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

OUTPUT = "data_dictionary.xlsx"

# ── Sheet 1: Data Dictionary df_final_full ─────────────────────────────────────
df_dict = pd.DataFrame([
    (
        "file_path",
        "String (object)",
        "Path lengkap ke file gambar",
        "/content/pothole_dataset/archive/.../train/images/pic-41.jpg\n"
        "/content/augmented_pothole_images/augmented_0.jpg",
        "Tidak",
        "Unik per baris. Path berbeda antara data Original (folder YOLO) "
        "dan Augmented (folder augmented_pothole_images)."
    ),
    (
        "label",
        "String (object)",
        "Nama folder induk tempat gambar berada — digunakan sebagai kelas/kategori gambar",
        "'images'  → gambar dari dataset YOLO original (train/images/ atau valid/images/)\n"
        "'augmented_pothole_images'  → gambar hasil augmentasi Albumentations",
        "Tidak",
        "Nilai 'images' muncul karena os.path.basename(root) mengambil "
        "nama folder terakhir. Semua gambar original berada di subfolder images/, "
        "sehingga labelnya 'images'. Augmented berada langsung di folder "
        "augmented_pothole_images/ sehingga labelnya nama folder itu."
    ),
    (
        "source",
        "String (object)",
        "Asal-usul gambar — apakah dari dataset original Roboflow atau hasil augmentasi",
        "'Original'   → dari dataset Pothole_Segmentation_YOLOv8 (Roboflow)\n"
        "'Augmented'  → dibuat dengan Albumentations dari gambar pothole",
        "Tidak",
        "Diisi secara manual saat memanggil collect_all_data(path, source_type). "
        "Total: 780 Original + 30 Augmented = 810 baris."
    ),
    (
        "subset",
        "String (object)",
        "Pembagian dataset untuk keperluan training machine learning",
        "'train'  → 80% data (648 gambar: 624 original + 24 augmented)\n"
        "'valid'  → 10% data (81 gambar: 78 original + 3 augmented)\n"
        "'test'   → 10% data (81 gambar: 78 original + 3 augmented)",
        "Tidak",
        "Dibagi dengan train_test_split (random_state=42, stratify='label'). "
        "Kolom ini ditambahkan setelah splitting, tidak ada di df_total."
    ),
], columns=[
    "Nama Kolom",
    "Tipe Data",
    "Deskripsi",
    "Contoh Nilai",
    "Boleh Null?",
    "Catatan Tambahan",
])

# ── Sheet 2: Ringkasan Statistik ───────────────────────────────────────────────
df_stats = pd.DataFrame([
    ("Total baris (gambar)",   "810",                      "780 Original + 30 Augmented"),
    ("Jumlah kolom",           "4",                        "file_path, label, source, subset"),
    ("Nilai null",             "0 (tidak ada)",            "DataFrame bersih, tidak ada missing value"),
    ("Duplikat",               "0 (tidak ada)",            "Setiap file_path unik"),
    ("Distribusi label",       "images: 780 | augmented_pothole_images: 30",
                                                           "Semua gambar adalah pothole (tidak ada kelas 'normal' di dataset ini)"),
    ("Distribusi source",      "Original: 780 | Augmented: 30",
                                                           "Rasio 96.3% : 3.7%"),
    ("Distribusi subset train","648 baris (80.0%)",        "624 original + 24 augmented"),
    ("Distribusi subset valid", "81 baris (10.0%)",        "78 original + 3 augmented"),
    ("Distribusi subset test",  "81 baris (10.0%)",        "78 original + 3 augmented"),
    ("File CSV output",        "data_pothole_gabungan_full.csv", "Disimpan dari df_final_full.to_csv()"),
], columns=["Atribut", "Nilai", "Keterangan"])

# ── Sheet 3: Proses Pembentukan df_final_full ─────────────────────────────────
df_proses = pd.DataFrame([
    ("1", "collect_all_data(path_original, 'Original')",
     "Scan rekursif folder pothole_dataset, ambil semua file JPG/PNG. "
     "Label = nama folder induk file (os.path.basename(root))."),
    ("2", "collect_all_data(path_augmented, 'Augmented')",
     "Scan folder augmented_pothole_images, ambil 30 file JPG. "
     "Label = 'augmented_pothole_images'."),
    ("3", "df_total = pd.DataFrame(all_orig + all_aug)",
     "Gabungkan semua data menjadi satu DataFrame. Total 810 baris, 3 kolom "
     "(file_path, label, source). Belum ada kolom subset."),
    ("4", "train_test_split(df_total, test_size=0.2, stratify='label')",
     "Split 80/20 dengan stratifikasi per label agar proporsi terjaga. "
     "Menghasilkan train_df (648 baris) dan temp_df (162 baris)."),
    ("5", "train_test_split(temp_df, test_size=0.5, stratify='label')",
     "Split temp_df 50/50 menjadi valid_df (81 baris) dan test_df (81 baris)."),
    ("6", "train_df['subset'] = 'train' (dst.)",
     "Tambahkan kolom subset ke masing-masing DataFrame hasil split."),
    ("7", "df_final_full = pd.concat([train_df, valid_df, test_df])",
     "Gabungkan kembali dengan ignore_index=True. "
     "Hasil akhir: 810 baris, 4 kolom (file_path, label, source, subset)."),
    ("8", "df_final_full.to_csv('data_pothole_gabungan_full.csv')",
     "Simpan ke file CSV sebagai output final tahap Data Cleaning."),
], columns=["Langkah", "Kode / Fungsi", "Penjelasan"])

# ── Tulis ke Excel ─────────────────────────────────────────────────────────────
sheets = {
    "Data Dictionary":          df_dict,
    "Ringkasan Statistik":      df_stats,
    "Proses Pembentukan":       df_proses,
}

with pd.ExcelWriter(OUTPUT, engine="openpyxl") as writer:
    for name, df in sheets.items():
        df.to_excel(writer, sheet_name=name, index=False)

# ── Styling ────────────────────────────────────────────────────────────────────
wb = load_workbook(OUTPUT)

HEADER_FILL = PatternFill("solid", fgColor="1A3A5C")
HEADER_FONT = Font(bold=True, color="F0A500", name="Calibri", size=11)
EVEN_FILL   = PatternFill("solid", fgColor="1E2A3A")
ODD_FILL    = PatternFill("solid", fgColor="161B27")
CELL_FONT   = Font(color="E0E0E0", name="Calibri", size=10)
BORDER      = Border(
    left=Side(style="thin", color="2D3A5A"),
    right=Side(style="thin", color="2D3A5A"),
    top=Side(style="thin", color="2D3A5A"),
    bottom=Side(style="thin", color="2D3A5A"),
)

COL_WIDTHS = {
    "Data Dictionary":     [18, 18, 32, 52, 12, 52],
    "Ringkasan Statistik": [28, 36, 45],
    "Proses Pembentukan":  [10, 42, 60],
}

for ws in wb.worksheets:
    # Header
    for cell in ws[1]:
        cell.fill      = HEADER_FILL
        cell.font      = HEADER_FONT
        cell.border    = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Data rows
    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=1):
        fill = EVEN_FILL if row_idx % 2 == 0 else ODD_FILL
        for cell in row:
            cell.fill      = fill
            cell.font      = CELL_FONT
            cell.border    = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    # Column widths
    widths = COL_WIDTHS.get(ws.title, [])
    for i, col in enumerate(ws.columns):
        letter = col[0].column_letter
        ws.column_dimensions[letter].width = widths[i] if i < len(widths) else 20

    # Row heights
    ws.row_dimensions[1].height = 24
    for i in range(2, ws.max_row + 1):
        ws.row_dimensions[i].height = 60 if ws.title == "Data Dictionary" else 30

    ws.freeze_panes = "A2"
    ws.sheet_properties.tabColor = "F0A500"

wb.save(OUTPUT)
print(f"[OK] {OUTPUT} berhasil dibuat.")
print(f"     Sheets: {', '.join(sheets.keys())}")
